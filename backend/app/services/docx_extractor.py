from docx import Document


def extract_text_from_docx(file_path: str) -> str:
    document = Document(file_path)

    extracted_parts = []

    # Extract normal paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            extracted_parts.append(text)

    # Extract text from tables
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                extracted_parts.append(" | ".join(row_text))

    return "\n".join(extracted_parts).strip()